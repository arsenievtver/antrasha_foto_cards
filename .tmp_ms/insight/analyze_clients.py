"""Итерация 2: профили клиентов, клиенты уровня A на контроле, подбор со склада.

Читает data/*.json (fetch_base.py), пишет:
  iter2_result.json          — агрегаты без ФИО (для анализа)
  clients_watchlist.xlsx     — список на обзвон с подборкой (ФИО, в git не идёт)
"""

from __future__ import annotations

import json
import re
import statistics
import time
from collections import Counter, defaultdict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ms_common import brand_key, canon_category
from ms_mcp import HERE, MsMcp

DATA = HERE / "data"
TODAY = date(2026, 10, 27)
CUR_FROM = date(2025, 10, 27)
PREV_FROM = date(2024, 10, 27)
PLACEHOLDER = re.compile(r"^(отказался|покупатель|розничный покупатель)$", re.I)
ARRIVAL = re.compile(r"/(\d{2})\.(\d{2})(?:\s|$)")
CARD_DISCOUNT_MAX = 20
OFF_SEASON = {"Бриджи, шорты"} if TODAY.month in (9, 10, 11, 12, 1, 2, 3) else set()
NON_SIZED = {"Аксессуары", "Бижутерия", "Сумка", "Ремни", "Галстук", "Бабочка", "Кошелек", "Очки", "Зонт"}

catalog = json.loads((DATA / "catalog.json").read_text())
for c in catalog.values():
    c["category"] = canon_category(c.get("path") or "", c.get("category"))
lines = json.loads((DATA / "sales_lines.json").read_text())
stock = json.loads((DATA / "stock_variants.json").read_text())


def arrival(article: str | None, name: str | None) -> date | None:
    mm = ARRIVAL.search(f"{article or ''} ") or ARRIVAL.search(f"{name or ''} ")
    if not mm:
        return None
    month, yy = int(mm.group(1)), int(mm.group(2))
    if not 1 <= month <= 12 or yy > TODAY.year % 100:
        return None
    return date(2000 + yy, month, 1)


def months_between(a: date, b: date) -> int:
    return (b.year - a.year) * 12 + b.month - a.month


for line in lines:
    c = catalog.get(line["product_id"]) or {}
    line["day"] = datetime.strptime(line["moment"][:10], "%Y-%m-%d").date()
    line["gender"] = c.get("gender")
    line["category"] = c.get("category")
    line["brand"] = brand_key(c.get("brand"))
    arr = arrival(c.get("article"), c.get("name"))
    line["age_at_sale_m"] = months_between(arr, line["day"]) if arr else None

out: dict = {}

# ---------- скидки и новинки (без ФИО) ----------


def bucket_age(m: int | None) -> str:
    if m is None:
        return "без даты"
    if m <= 3:
        return "0-3 мес"
    if m <= 6:
        return "4-6 мес"
    if m <= 12:
        return "7-12 мес"
    if m <= 24:
        return "13-24 мес"
    return "25+ мес"


age_acc = defaultdict(lambda: {"units": 0, "revenue": 0.0, "list": 0.0, "markdown_units": 0})
for line in lines:
    a = age_acc[bucket_age(line["age_at_sale_m"])]
    a["units"] += line["qty"]
    a["revenue"] += line["sum"]
    a["list"] += line["qty"] * line["list_price"]
    if line["discount_pct"] > CARD_DISCOUNT_MAX:
        a["markdown_units"] += line["qty"]
total_rev = sum(v["revenue"] for v in age_acc.values())
out["sales_by_age_at_sale_24m"] = {
    k: {
        "units": v["units"],
        "revenue_share_pct": round(v["revenue"] / total_rev * 100, 1),
        "avg_discount_pct": round((1 - v["revenue"] / v["list"]) * 100, 1) if v["list"] else None,
        "markdown_units_pct": round(v["markdown_units"] / v["units"] * 100, 1) if v["units"] else None,
    }
    for k, v in sorted(age_acc.items())
}

# ---------- профили клиентов ----------

clients: dict[str, dict] = {}
for line in lines:
    if not line["agent_id"] or PLACEHOLDER.match((line["agent"] or "").strip()):
        continue
    cl = clients.setdefault(line["agent_id"], {"name": line["agent"], "lines": []})
    cl["lines"].append(line)

profiles = []
for cid, cl in clients.items():
    ls = [l for l in cl["lines"] if l["day"] >= PREV_FROM]
    if not ls:
        continue
    days = sorted({l["day"] for l in ls})
    gaps = [(b - a).days for a, b in zip(days, days[1:])]
    rev24 = sum(l["sum"] for l in ls)
    rev_cur = sum(l["sum"] for l in ls if l["day"] >= CUR_FROM)
    rev_prev = sum(l["sum"] for l in ls if PREV_FROM <= l["day"] < CUR_FROM)
    items = sum(l["qty"] for l in ls)
    genders = Counter(l["gender"] for l in ls if l["gender"] in ("M", "W"))
    brands = Counter(l["brand"] for l in ls if l["brand"])
    cats = Counter(l["category"] for l in ls if l["category"])
    sizes_by_cat: dict[str, Counter] = defaultdict(Counter)
    for l in ls:
        if l["size"] and l["category"] and l["category"] not in NON_SIZED:
            sizes_by_cat[(l["gender"], l["category"])][l["size"]] += l["qty"]
    listed = sum(l["qty"] * l["list_price"] for l in ls)
    markdown_units = sum(l["qty"] for l in ls if l["discount_pct"] > CARD_DISCOUNT_MAX)
    top_brand, top_brand_n = (brands.most_common(1) or [(None, 0)])[0]
    profiles.append(
        {
            "id": cid,
            "name": cl["name"],
            "rev24": rev24,
            "rev_cur": rev_cur,
            "rev_prev": rev_prev,
            "visits": len(days),
            "items": items,
            "first": days[0],
            "last": days[-1],
            "days_since": (TODAY - days[-1]).days,
            "median_gap": statistics.median(gaps) if gaps else None,
            "gender": genders.most_common(1)[0][0] if genders else None,
            "brands": brands,
            "top_brand": top_brand,
            "top_brand_share": top_brand_n / sum(brands.values()) if brands else 0,
            "cats": cats,
            "sizes": {k: v.most_common(1)[0][0] for k, v in sizes_by_cat.items()},
            "avg_discount": (1 - rev24 / listed) * 100 if listed else 0,
            "markdown_share": markdown_units / items if items else 0,
            "bought_products": {l["product_id"] for l in ls},
        }
    )

profiles.sort(key=lambda p: -p["rev24"])
n_a = max(1, round(len(profiles) * 0.2))
for i, p in enumerate(profiles):
    p["segment"] = "A" if i < n_a else ("B" if i < n_a * 2.5 else "C")


def status(p: dict) -> str:
    gap = p["median_gap"]
    if p["days_since"] > 365:
        return "ушёл (>12 мес)"
    if gap is None:
        return "одна покупка"
    limit = max(1.5 * gap, 90)
    if p["days_since"] > limit:
        return "просрочил визит"
    if p["rev_prev"] > 0 and p["rev_cur"] < 0.5 * p["rev_prev"]:
        return "покупает меньше"
    return "в ритме"


for p in profiles:
    p["status"] = status(p)

seg_acc = defaultdict(lambda: {"clients": 0, "rev24": 0.0, "rev_cur": 0.0, "rev_prev": 0.0, "markdown_units": 0.0, "items": 0})
for p in profiles:
    s = seg_acc[p["segment"]]
    s["clients"] += 1
    s["rev24"] += p["rev24"]
    s["rev_cur"] += p["rev_cur"]
    s["rev_prev"] += p["rev_prev"]
    s["markdown_units"] += p["markdown_share"] * p["items"]
    s["items"] += p["items"]
out["segments"] = {
    k: {
        "clients": v["clients"],
        "rev24": round(v["rev24"]),
        "rev_prev12": round(v["rev_prev"]),
        "rev_cur12": round(v["rev_cur"]),
        "markdown_units_pct": round(v["markdown_units"] / v["items"] * 100, 1) if v["items"] else None,
    }
    for k, v in sorted(seg_acc.items())
}
out["a_status"] = {}
for st in sorted({p["status"] for p in profiles}):
    group = [p for p in profiles if p["segment"] == "A" and p["status"] == st]
    if group:
        out["a_status"][st] = {
            "clients": len(group),
            "rev_prev12": round(sum(p["rev_prev"] for p in group)),
            "rev_cur12": round(sum(p["rev_cur"] for p in group)),
        }

# ---------- гипотеза «бренд или категория/размер» ----------

loyal_pool = [p for p in profiles if p["items"] >= 5 and p["brands"]]
by_gender = defaultdict(lambda: {"clients": 0, "brand_loyal": 0, "top_brand_share_sum": 0.0})
for p in loyal_pool:
    g = by_gender[p["gender"] or "-"]
    g["clients"] += 1
    g["top_brand_share_sum"] += p["top_brand_share"]
    if p["top_brand_share"] >= 0.6:
        g["brand_loyal"] += 1
out["brand_loyalty_clients_5plus_items"] = {
    k: {
        "clients": v["clients"],
        "brand_loyal_ge60pct": v["brand_loyal"],
        "avg_top_brand_share_pct": round(v["top_brand_share_sum"] / v["clients"] * 100, 1),
    }
    for k, v in by_gender.items()
}
loyal_brands = Counter(p["top_brand"] for p in loyal_pool if p["top_brand_share"] >= 0.6)
out["brands_with_loyal_clients"] = loyal_brands.most_common(15)

# Повторная покупка той же категории: тот же бренд или другой
same = diff = 0
for p in profiles:
    seen: dict[tuple, set] = defaultdict(set)
    for l in sorted(clients[p["id"]]["lines"], key=lambda x: x["day"]):
        key = (l["gender"], l["category"])
        if l["brand"] and l["category"]:
            if seen[key]:
                if l["brand"] in seen[key]:
                    same += 1
                else:
                    diff += 1
            seen[key].add(l["brand"])
out["repeat_category_purchase"] = {
    "same_brand": same,
    "other_brand": diff,
    "same_brand_pct": round(same / (same + diff) * 100, 1) if same + diff else None,
}

# ---------- подбор со склада для A на контроле ----------

stock_by_key = defaultdict(list)
for s in stock:
    c = catalog.get(s["product_id"]) or {}
    if not c or not s["size"] or c.get("top") in ("Сертификат", "Сертификаты", "Товар ТО"):
        continue
    arr = arrival(c.get("article"), c.get("name"))
    s["age_m"] = months_between(arr, TODAY) if arr else None
    s["brand"] = brand_key(c.get("brand"))
    s["category"] = c.get("category")
    s["gender"] = c.get("gender")
    stock_by_key[(c.get("gender"), c.get("category"), s["size"])].append(s)

watch = [p for p in profiles if p["segment"] == "A" and p["status"] in ("просрочил визит", "покупает меньше", "ушёл (>12 мес)")]
for p in watch:
    picks = []
    top_cats = [c for c, _ in p["cats"].most_common(6)]
    for (g, cat), size in p["sizes"].items():
        if cat in OFF_SEASON:
            continue
        rank = top_cats.index(cat) if cat in top_cats else 6
        for s in stock_by_key.get((g, cat, size), []):
            if s["product_id"] in p["bought_products"]:
                continue
            brand_hit = s["brand"] in p["brands"]
            score = (3 if brand_hit else 0) + (6 - rank) * 0.5 + (1 if (s["age_m"] or 99) <= 6 else 0)
            picks.append((score, s, brand_hit))
    picks.sort(key=lambda t: -t[0])
    seen_products = set()
    per_cat: Counter = Counter()
    fresh, old = [], []
    for score, s, brand_hit in picks:
        if s["product_id"] in seen_products or per_cat[s["category"]] >= 2:
            continue
        age = s["age_m"]
        if age is not None and age <= 6 and len(fresh) < 4:
            fresh.append((s, brand_hit))
        elif age is not None and age > 12 and len(old) < 3:
            old.append((s, brand_hit))
        else:
            continue
        seen_products.add(s["product_id"])
        per_cat[s["category"]] += 1
    p["picks"] = fresh + old

out["watchlist"] = {
    "clients": len(watch),
    "with_picks": sum(1 for p in watch if p["picks"]),
    "picks_total": sum(len(p["picks"]) for p in watch),
    "picks_fresh_6m": sum(1 for p in watch for s, _ in p["picks"] if (s["age_m"] or 99) <= 6),
    "picks_older_12m": sum(1 for p in watch for s, _ in p["picks"] if (s["age_m"] or 0) > 12),
    "picks_brand_match": sum(1 for p in watch for _, b in p["picks"] if b),
    "picks_retail_value": round(sum(s["sale_price"] or 0 for p in watch for s, _ in p["picks"])),
    "rev_prev12": round(sum(p["rev_prev"] for p in watch)),
    "rev_cur12": round(sum(p["rev_cur"] for p in watch)),
}

# Сколько старого склада (>12 мес) можно предложить лично: есть клиент A/B
# с тем же полом, категорией, размером и этим брендом в истории.
demand_keys = defaultdict(set)
for p in profiles:
    if p["segment"] not in ("A", "B"):
        continue
    for (g, cat), size in p["sizes"].items():
        for b in p["brands"]:
            demand_keys[(g, cat, size, b)].add(p["id"])
old_total = {"units": 0, "cost": 0.0}
old_matched = {"units": 0, "cost": 0.0, "clients": set()}
for items in stock_by_key.values():
    for s in items:
        if (s["age_m"] or 0) <= 12:
            continue
        old_total["units"] += s["stock"]
        old_total["cost"] += s["stock"] * (s["cost"] or 0)
        who = demand_keys.get((s["gender"], s["category"], s["size"], s["brand"]))
        if who:
            old_matched["units"] += s["stock"]
            old_matched["cost"] += s["stock"] * (s["cost"] or 0)
            old_matched["clients"] |= who
out["old_stock_personal_match"] = {
    "old_sized_units": old_total["units"],
    "old_sized_cost": round(old_total["cost"]),
    "matched_units": old_matched["units"],
    "matched_cost": round(old_matched["cost"]),
    "matched_clients_ab": len(old_matched["clients"]),
}

# ---------- телефоны и xlsx ----------

m = MsMcp()
for p in watch:
    try:
        cp = m.call("counterparty_get", {"id": p["id"]})
        p["phone"] = cp.get("phone")
        time.sleep(0.5)
    except RuntimeError:
        p["phone"] = None

wb = Workbook()
ws = wb.active
ws.title = "Клиенты A на контроле"
head = [
    "Клиент", "Телефон", "Статус", "Последняя покупка", "Дней назад", "Обычный интервал, дн",
    "Выручка пред. 12 мес", "Выручка тек. 12 мес", "Визитов за 24 мес", "Пол", "Любимые бренды",
    "Категории", "Размеры", "Средняя скидка, %", "Доля уценки, %", "Подборка со склада",
]
ws.append(head)
order = {"просрочил визит": 0, "покупает меньше": 1, "ушёл (>12 мес)": 2}
for p in sorted(watch, key=lambda p: (order[p["status"]], -p["rev24"])):
    picks_txt = "\n".join(
        f'{s["name"]} — ' + f'{round(s["sale_price"] or 0):,} ₽'.replace(",", " ")
        + (f' [{s["age_m"]} мес]' if s.get("age_m") is not None else "")
        + (" ★бренд" if b else "")
        for s, b in p["picks"]
    )
    ws.append([
        p["name"], p.get("phone"), p["status"], p["last"].isoformat(), p["days_since"],
        round(p["median_gap"]) if p["median_gap"] else None,
        round(p["rev_prev"]), round(p["rev_cur"]), p["visits"],
        {"M": "муж", "W": "жен"}.get(p["gender"], ""),
        ", ".join(b for b, _ in p["brands"].most_common(3)),
        ", ".join(c for c, _ in p["cats"].most_common(4)),
        "; ".join(f"{c}: {s}" for (g, c), s in p["sizes"].items()),
        round(p["avg_discount"]), round(p["markdown_share"] * 100), picks_txt,
    ])
widths = [28, 16, 16, 14, 10, 12, 14, 14, 10, 6, 30, 34, 40, 10, 10, 70]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for cell in ws[1]:
    cell.font = Font(bold=True)
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "B2"
wb.save(HERE / "clients_watchlist.xlsx")

(HERE / "iter2_result.json").write_text(json.dumps(out, ensure_ascii=False, indent=1, default=str))
print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
