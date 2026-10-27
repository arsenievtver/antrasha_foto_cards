"""Итерация 3: ассортимент — распродажа по сезонам, размерная сетка, бренды, хиты и неликвид.

Читает data/catalog.json, sales_lines.json, stock_variants.json, arrivals.json, losses.json.
Пишет iter3_result.json (агрегаты) и assortment_report.xlsx.
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ms_common import base_name, brand_key, canon_category
from ms_mcp import HERE, MsMcp

DATA = HERE / "data"
TODAY = date(2026, 10, 27)
CARD_DISCOUNT_MAX = 20
FAST_SELLOUT_DAYS = 45
SIZE_WINDOW_DAYS = 120
MAIN_SEASONS = ["ОЗ2024", "ВЛ2025", "ОЗ2025", "ВЛ2026", "ОЗ2026"]
CURVE_SEASONS = {"ВЛ2025", "ОЗ2025", "ВЛ2026"}
SEASON_RE = re.compile(r"(ВЛ|ОЗ)\s*(20\d\d)")
LETTERS = {"XXS": 0, "XS": 1, "S": 2, "M": 3, "L": 4, "XL": 5, "XXL": 6, "2XL": 6, "XXXL": 7, "3XL": 7}


def day_of(moment: str) -> date:
    return datetime.strptime(moment[:10], "%Y-%m-%d").date()


def size_sort(s: str) -> tuple:
    m = re.match(r"(\d+)", s)
    if m:
        return (0, int(m.group(1)), s)
    return (1, LETTERS.get(s.upper(), 99), s)


def season_of(name: str, first: date) -> str:
    m = SEASON_RE.search(name or "")
    if m:
        tag = f"{m.group(1)}{m.group(2)}"
        # ОЗ2027 в названиях — это текущий осенне-зимний сезон 2026 года
        return "ОЗ2026" if tag == "ОЗ2027" else tag
    if first.month == 12:
        return f"ВЛ{first.year + 1}"
    if first.month <= 5:
        return f"ВЛ{first.year}"
    return f"ОЗ{first.year}"


catalog = json.loads((DATA / "catalog.json").read_text())
by_name: dict[str, str] = {}
for pid, c in catalog.items():
    c["category"] = canon_category(c.get("path") or "", c.get("category"))
    c["brand"] = brand_key(c.get("brand"))
    by_name.setdefault((c["name"] or "").strip(), pid)

arrivals = json.loads((DATA / "arrivals.json").read_text())
sales = json.loads((DATA / "sales_lines.json").read_text())
stock = json.loads((DATA / "stock_variants.json").read_text())
losses = json.loads((DATA / "losses.json").read_text())

variants: dict[str, dict] = {}
for a in arrivals:
    key = (a["variant_name"] or "").strip()
    if not key:
        continue
    d = day_of(a["moment"])
    v = variants.setdefault(
        key,
        {"name": key, "received": 0, "cost_sum": 0.0, "first": d, "size": a["size"], "sales": [], "lost": 0, "stock": 0},
    )
    v["received"] += a["qty"]
    v["cost_sum"] += a["qty"] * a["cost"]
    v["first"] = min(v["first"], d)

for s in sales:
    v = variants.get((s["variant_name"] or "").strip())
    if v:
        v["sales"].append(s)
for s in stock:
    v = variants.get((s["name"] or "").strip())
    if v:
        v["stock"] += s["stock"] or 0
for l in losses:
    v = variants.get((l["variant_name"] or "").strip())
    if v:
        v["lost"] += l["qty"]

rows = []
for v in variants.values():
    pid = by_name.get(base_name(v["name"]))
    c = catalog.get(pid) or {}
    if c.get("top") in ("Сертификат", "Сертификаты", "Товар ТО") or not c:
        continue
    first = v["first"]
    if any(day_of(s["moment"]) < first for s in v["sales"]):
        continue  # модель была на складе до окна выгрузки — когорта неполная
    unit_cost = v["cost_sum"] / v["received"] if v["received"] else 0
    cum = 0
    sellout_day = None
    for s in sorted(v["sales"], key=lambda x: x["moment"]):
        cum += s["qty"]
        if sellout_day is None and cum >= v["received"]:
            sellout_day = (day_of(s["moment"]) - first).days

    def sold_within(days: int, card_only: bool = False) -> float:
        return sum(
            s["qty"]
            for s in v["sales"]
            if (day_of(s["moment"]) - first).days <= days and (not card_only or s["discount_pct"] <= CARD_DISCOUNT_MAX)
        )

    rows.append(
        {
            "name": v["name"],
            "product_id": pid,
            "model": c.get("name"),
            "brand": c.get("brand") or "—",
            "gender": c.get("gender"),
            "category": c.get("category"),
            "size": v["size"],
            "season": season_of(v["name"], first),
            "first": first,
            "age": (TODAY - first).days,
            "received": v["received"],
            "unit_cost": unit_cost,
            "sold_total": sum(s["qty"] for s in v["sales"]),
            "sold_card_total": sum(s["qty"] for s in v["sales"] if s["discount_pct"] <= CARD_DISCOUNT_MAX),
            "revenue": sum(s["sum"] for s in v["sales"]),
            "sold90": sold_within(90),
            "sold180": sold_within(180),
            "sold_size_window": sold_within(SIZE_WINDOW_DAYS),
            "card180": sold_within(180, card_only=True),
            "sellout_day": sellout_day,
            "stock": v["stock"],
            "lost": v["lost"],
        }
    )


def agg(items: list[dict]) -> dict:
    rec = sum(r["received"] for r in items)
    cost = sum(r["received"] * r["unit_cost"] for r in items)
    e90 = [r for r in items if r["age"] >= 90]
    e180 = [r for r in items if r["age"] >= 180]
    rec90 = sum(r["received"] for r in e90)
    rec180 = sum(r["received"] for r in e180)
    sold = sum(r["sold_total"] for r in items)
    revenue = sum(r["revenue"] for r in items)
    cogs = sum(r["sold_total"] * r["unit_cost"] for r in items)
    return {
        "models": len({r["product_id"] for r in items}),
        "received": rec,
        "cost_received": round(cost),
        "st90_pct": round(sum(r["sold90"] for r in e90) / rec90 * 100, 1) if rec90 else None,
        "st180_pct": round(sum(r["sold180"] for r in e180) / rec180 * 100, 1) if rec180 else None,
        "st180_card_pct": round(sum(r["card180"] for r in e180) / rec180 * 100, 1) if rec180 else None,
        "st_to_date_pct": round(sold / rec * 100, 1) if rec else None,
        "card_share_of_sold_pct": round(sum(r["sold_card_total"] for r in items) / sold * 100, 1) if sold else None,
        "revenue": round(revenue),
        "gross_profit": round(revenue - cogs),
        "revenue_per_rub_bought": round(revenue / cost, 2) if cost else None,
        "stock_units": sum(r["stock"] for r in items),
        "stock_cost": round(sum(r["stock"] * r["unit_cost"] for r in items)),
    }


def group(key_fn, items: list[dict]) -> dict:
    g = defaultdict(list)
    for r in items:
        g[key_fn(r)].append(r)
    return g


out: dict = {}
main = [r for r in rows if r["season"] in MAIN_SEASONS]
out["seasons"] = {s: agg(items) for s, items in sorted(group(lambda r: r["season"], main).items(), key=lambda kv: MAIN_SEASONS.index(kv[0]))}

g_label = {"M": "муж", "W": "жен"}
curve = [r for r in main if r["season"] in CURVE_SEASONS]
cat_rows = []
for (g, cat), items in group(lambda r: (r["gender"], r["category"]), curve).items():
    a = agg(items)
    if a["received"] >= 30:
        cat_rows.append({"gender": g_label.get(g, "-"), "category": cat, **a})
cat_rows.sort(key=lambda r: -r["cost_received"])
out["categories_vl25_oz25_vl26"] = cat_rows

# Размерная сетка: доля прихода и доля продаж за первые 120 дней + быстрый распродажа размера
size_rows = []
for (g, cat), items in group(lambda r: (r["gender"], r["category"]), [r for r in curve if r["size"] and r["age"] >= SIZE_WINDOW_DAYS]).items():
    rec_total = sum(r["received"] for r in items)
    sold_total = sum(r["sold_size_window"] for r in items)
    if rec_total < 40 or sold_total < 15:
        continue
    for size, its in sorted(group(lambda r: r["size"], items).items(), key=lambda kv: size_sort(kv[0])):
        rec = sum(r["received"] for r in its)
        sold = sum(r["sold_size_window"] for r in its)
        fast = sum(1 for r in its if r["sellout_day"] is not None and r["sellout_day"] <= FAST_SELLOUT_DAYS)
        rec_share = rec / rec_total
        sold_share = sold / sold_total if sold_total else 0
        index = sold_share / rec_share if rec_share else None
        fast_pct = fast / len(its)
        if rec < 3:
            verdict = ""
        elif (index or 0) >= 1.25 or (fast_pct >= 0.35 and (index or 0) >= 1):
            verdict = "добавить"
        elif (index or 0) <= 0.75 and fast_pct < 0.15:
            verdict = "сократить"
        else:
            verdict = ""
        size_rows.append(
            {
                "gender": g_label.get(g, "-"),
                "category": cat,
                "size": size,
                "received": rec,
                "received_share_pct": round(rec_share * 100, 1),
                "sold120": sold,
                "sold_share_pct": round(sold_share * 100, 1),
                "index": round(index, 2) if index is not None else None,
                "model_sizes": len(its),
                "fast_sellout_pct": round(fast_pct * 100),
                "verdict": verdict,
            }
        )
out["size_curve_flags"] = [r for r in size_rows if r["verdict"]]
fast_all = [r for r in curve if r["size"] and r["age"] >= SIZE_WINDOW_DAYS]
out["fast_sellout_overall"] = {
    "model_sizes": len(fast_all),
    "sold_out_within_45d": sum(1 for r in fast_all if r["sellout_day"] is not None and r["sellout_day"] <= FAST_SELLOUT_DAYS),
    "single_unit_share_pct": round(sum(1 for r in fast_all if r["received"] == 1) / len(fast_all) * 100, 1) if fast_all else None,
}

brand_rows = []
for (b, g, cat), items in group(lambda r: (r["brand"], r["gender"], r["category"]), curve).items():
    a = agg(items)
    if a["received"] >= 15:
        brand_rows.append({"brand": b, "gender": g_label.get(g, "-"), "category": cat, **a})
brand_rows.sort(key=lambda r: (r["gender"], r["category"], -(r["revenue_per_rub_bought"] or 0)))
brand_total = []
for b, items in group(lambda r: r["brand"], curve).items():
    a = agg(items)
    if a["received"] >= 30:
        brand_total.append({"brand": b, **a})
brand_total.sort(key=lambda r: -r["cost_received"])
out["brands_total_vl25_oz25_vl26"] = brand_total

model_rows = []
for pid, items in group(lambda r: r["product_id"], main).items():
    a = agg(items)
    r0 = items[0]
    model_rows.append(
        {
            "model": r0["model"],
            "brand": r0["brand"],
            "gender": g_label.get(r0["gender"], "-"),
            "category": r0["category"],
            "season": r0["season"],
            "first": min(r["first"] for r in items).isoformat(),
            "sizes": ", ".join(sorted({r["size"] for r in items if r["size"]}, key=size_sort)),
            "sold_out_sizes": ", ".join(sorted({r["size"] for r in items if r["size"] and r["stock"] <= 0 and r["sold_total"] > 0}, key=size_sort)),
            "age_days": min(r["age"] for r in items),
            **a,
        }
    )
hits = [m for m in model_rows if m["received"] >= 4 and m["st90_pct"] is not None]
hits.sort(key=lambda m: (-(m["st90_pct"] or 0), -(m["card_share_of_sold_pct"] or 0)))
dead = [m for m in model_rows if m["received"] >= 3 and m["st180_pct"] is not None and m["st180_pct"] <= 20]
dead.sort(key=lambda m: -m["stock_cost"])
out["models"] = {
    "total": len(model_rows),
    "hits_st90_ge70": sum(1 for m in hits if m["st90_pct"] >= 70),
    "dead_st180_le20": len(dead),
    "dead_stock_cost": round(sum(m["stock_cost"] for m in dead)),
    "dead_cost_received": round(sum(m["cost_received"] for m in dead)),
}
dead_by_brand = defaultdict(lambda: {"models": 0, "stock_cost": 0})
for m in dead:
    dead_by_brand[m["brand"]]["models"] += 1
    dead_by_brand[m["brand"]]["stock_cost"] += m["stock_cost"]
out["dead_by_brand"] = sorted(([b, v["models"], v["stock_cost"]] for b, v in dead_by_brand.items()), key=lambda x: -x[2])[:12]

# Закупки против себестоимости продаж и рост склада
mcp = MsMcp()


def stock_cost_on(moment: str) -> float:
    rows_ = mcp.call_all("report_stock_all", {"stockMode": "positiveOnly", "groupBy": "product", "moment": moment})
    return sum((r.get("stock") or 0) * (r.get("price") or 0) for r in rows_)


def cogs(date_from: str, date_to: str) -> float:
    d = mcp.call("report_profit_byproduct", {"momentFrom": date_from, "momentTo": date_to, "limit": 1, "includeTotals": True})
    t = d.get("totals") or {}
    return (t.get("sellCostSum") or 0) - (t.get("returnCostSum") or 0)


def bought(date_from: date, date_to: date) -> float:
    return sum(a["qty"] * a["cost"] for a in arrivals if date_from <= day_of(a["moment"]) <= date_to)


out["stock_flow"] = {
    "stock_cost_2024-10-26": round(stock_cost_on("2024-10-26")),
    "stock_cost_2025-10-26": round(stock_cost_on("2025-10-26")),
    "stock_cost_now": round(stock_cost_on("2026-10-27")),
    "bought_cost_2024-10-27..2025-10-26": round(bought(date(2024, 10, 27), date(2025, 10, 26))),
    "cogs_2024-10-27..2025-10-26": round(cogs("2024-10-27 00:00:00", "2025-10-26")),
    "bought_cost_2025-10-27..2026-10-26": round(bought(date(2025, 10, 27), date(2026, 10, 26))),
    "cogs_2025-10-27..2026-10-26": round(cogs("2025-10-27 00:00:00", "2026-10-26")),
}

(HERE / "iter3_result.json").write_text(json.dumps(out, ensure_ascii=False, indent=1, default=str))

# ---------- xlsx ----------
COLS = {
    "models": "Моделей", "received": "Пришло, шт", "cost_received": "Закуплено, ₽ себест.",
    "st90_pct": "Продано за 90 дн, %", "st180_pct": "Продано за 180 дн, %", "st180_card_pct": "Продано за 180 дн по карте, %",
    "st_to_date_pct": "Продано всего, %", "card_share_of_sold_pct": "Доля продаж по карте, %", "revenue": "Выручка, ₽",
    "gross_profit": "Валовая прибыль, ₽", "revenue_per_rub_bought": "Выручка на 1 ₽ закупки", "stock_units": "Остаток, шт",
    "stock_cost": "Остаток, ₽ себест.", "gender": "Пол", "category": "Категория", "brand": "Бренд", "size": "Размер",
    "received_share_pct": "Доля прихода, %", "sold120": "Продано за 120 дн, шт", "sold_share_pct": "Доля продаж, %",
    "index": "Индекс спроса", "model_sizes": "Моделе-размеров", "fast_sellout_pct": "Кончился за 45 дн, %",
    "verdict": "Вывод", "model": "Модель", "season": "Сезон", "first": "Первый приход", "sizes": "Размеры",
    "sold_out_sizes": "Распроданы размеры", "age_days": "Дней с прихода",
}


def sheet(wb: Workbook, title: str, data: list[dict], keys: list[str]) -> None:
    ws = wb.create_sheet(title)
    ws.append([COLS.get(k, k) for k in keys])
    for r in data:
        ws.append([r.get(k) for k in keys])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, k in enumerate(keys, 1):
        ws.column_dimensions[get_column_letter(i)].width = 34 if k in ("model", "sizes") else 16
    ws.freeze_panes = "A2"


metric_keys = ["models", "received", "cost_received", "st90_pct", "st180_pct", "st180_card_pct", "st_to_date_pct",
               "card_share_of_sold_pct", "revenue", "gross_profit", "revenue_per_rub_bought", "stock_units", "stock_cost"]
wb = Workbook()
wb.remove(wb.active)
sheet(wb, "Сезоны", [{"season": s, **a} for s, a in out["seasons"].items()], ["season"] + metric_keys)
sheet(wb, "Категории", cat_rows, ["gender", "category"] + metric_keys)
sheet(wb, "Размерная сетка", size_rows, ["gender", "category", "size", "received", "received_share_pct", "sold120",
                                          "sold_share_pct", "index", "model_sizes", "fast_sellout_pct", "verdict"])
sheet(wb, "Бренды в категории", brand_rows, ["gender", "category", "brand"] + metric_keys)
sheet(wb, "Бренды итого", brand_total, ["brand"] + metric_keys)
model_keys = ["model", "brand", "gender", "category", "season", "first", "age_days", "sizes", "sold_out_sizes"] + metric_keys[1:]
sheet(wb, "Хиты", hits[:150], model_keys)
sheet(wb, "Неликвид", dead, model_keys)
wb.save(HERE / "assortment_report.xlsx")

print(json.dumps({k: v for k, v in out.items() if k not in ("categories_vl25_oz25_vl26",)}, ensure_ascii=False, indent=1, default=str))
